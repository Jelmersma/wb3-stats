#!/usr/bin/env python3
"""Leest data/WB3 Stats.xlsx en schrijft site/assets/data.js.

Gebruik:  python3 build.py [pad/naar/excel.xlsx]

Bronnen per gegeven:
  - Minuten / wedstrijden  -> tabbladen Competitie en Oefen (per datum, op naam)
  - Trainingen             -> tabblad Training (per datum, op naam)
  - Goals, assists, kaarten, snuiter (totalen)
                           -> handmatige kolommen in tabblad Selectie
  - Per-wedstrijd details  -> tabblad Wedstrijden (uitslag + Stats-tekst)

De formules in Selectie (minuten, trainingen) worden bewust genegeerd: die
verwijzen op rijnummer en lopen scheef zodra de volgorde verschilt.
Alles wordt op naam gekoppeld, en de twee bronnen voor goals/assists/etc.
worden tegen elkaar gecontroleerd.
"""

import datetime as dt
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl ontbreekt: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent
DEFAULT_XLSX = ROOT / "data" / "WB3 Stats.xlsx"
OUT_JS = ROOT / "site" / "assets" / "data.js"
STAND_JSON = ROOT / "data" / "stand.json"
OUR_TEAM = "Woudenberg 3"  # zo heten wij in de competitiestand
SITE = ROOT / "site"

STAT_KEYS = ["goals", "assists", "snuiter", "geel", "rood"]
TEXT_KEYS = {
    "doelpunten": "goals", "doelpunt": "goals", "goals": "goals", "goal": "goals",
    "assists": "assists", "assist": "assists",
    "snuiter": "snuiter", "snuiters": "snuiter",
    "geel": "geel", "gele kaart": "geel", "gele kaarten": "geel",
    "rood": "rood", "rode kaart": "rood", "rode kaarten": "rood",
}
OWN_GOAL = {"eigen goal", "eigen doelpunt", "own goal", "og", "e.d.", "ed"}

# Andere namen die in de Stats-tekst gebruikt worden -> naam zoals in de selectie
ALIASES = {
    "robin h.": "Hanna",
    "robin h": "Hanna",
    "robin hanna": "Hanna",
}

warnings: list[str] = []
notes: list[str] = []


def warn(msg: str) -> None:
    warnings.append(msg)


# ---------------------------------------------------------------- helpers

def norm(name) -> str:
    return " ".join(str(name).split()).lower()


def squash(name) -> str:
    return re.sub(r"[^a-z0-9]", "", norm(name))


def slug(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def as_date(v) -> dt.datetime | None:
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime.combine(v, dt.time())
    return None


def sheet(wb, name: str, required=True):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    if required:
        sys.exit(f"FOUT: tabblad '{name}' niet gevonden in de Excel.")
    return None


def is_total(name) -> bool:
    return norm(name) in {"totaal", "total", "totalen"}


# ---------------------------------------------------------------- season / dates

class Season:
    """Corrigeert jaartallen: seizoen loopt van juli t/m juni."""

    def __init__(self, dates: list[dt.datetime]):
        autumn = [d for d in dates if d.month >= 7]
        self.start = min(d.year for d in autumn) if autumn else min(d.year for d in dates) - 1

    def fix(self, d: dt.datetime, where: str) -> dt.datetime:
        want = self.start if d.month >= 7 else self.start + 1
        if d.year != want:
            fixed = d.replace(year=want)
            notes.append(f"Jaartal gecorrigeerd in {where}: {d:%d-%m-%Y} -> {fixed:%d-%m-%Y}")
            return fixed
        return d

    @property
    def label(self) -> str:
        return f"{self.start}/{str(self.start + 1)[2:]}"


# ---------------------------------------------------------------- players

class Roster:
    def __init__(self):
        self.players: dict[str, dict] = {}   # norm name -> player
        self.order: list[str] = []

    def add(self, name: str, source: str) -> dict:
        key = norm(name)
        if key not in self.players:
            display = " ".join(str(name).split())
            pid = slug(display)
            self.players[key] = {"id": pid, "name": display, "sources": set()}
            self.order.append(key)
        p = self.players[key]
        p["sources"].add(source)
        return p

    def get(self, name) -> dict | None:
        return self.players.get(norm(name))

    def resolve(self, raw: str) -> tuple[dict | None, str]:
        """Zoek een speler bij een naam uit vrije tekst."""
        n = norm(raw)
        if n in ALIASES:
            n = norm(ALIASES[n])
        if n in self.players:
            return self.players[n], "exact"
        sq = squash(raw)
        hits = [p for k, p in self.players.items() if squash(k) == sq]
        if len(hits) == 1:
            return hits[0], "squash"
        first = [p for k, p in self.players.items() if k.split()[0] == n]
        if len(first) == 1:
            return first[0], "voornaam"
        if len(first) > 1:
            return None, "dubbel"
        return None, "onbekend"


# ---------------------------------------------------------------- grid sheets

def read_grid(ws, roster: Roster, season: Season, label: str):
    """Tabblad met namen in kolom A en datums als kolomkoppen.

    Geeft (sessions, rows) terug: sessions = [{date, col}] die gevuld zijn,
    rows = {norm name: {col: value}}.
    """
    header = [c.value for c in ws[1]]
    cols = []
    for idx, v in enumerate(header[1:], start=1):
        d = as_date(v)
        if d:
            cols.append((idx, season.fix(d, label)))

    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        name = r[0] if r else None
        if name is None or str(name).strip() == "":
            continue
        if is_total(name):
            break
        roster.add(name, label)
        rows[norm(name)] = {idx: num(r[idx]) if idx < len(r) else None for idx, _ in cols}

    sessions = []
    for idx, d in cols:
        if any(vals.get(idx) is not None for vals in rows.values()):
            sessions.append({"date": d, "col": idx})
    return sessions, rows


# ---------------------------------------------------------------- selectie (handmatige totalen)

def read_selectie(ws, roster: Roster):
    group_row = [c.value for c in ws[1]]
    head_row = [c.value for c in ws[2]]
    starts = []
    for i, v in enumerate(group_row):
        if v is None:
            continue
        low = str(v).lower()
        if low.startswith("comp"):
            starts.append((i, "comp"))
        elif low.startswith("oefen"):
            starts.append((i, "oefen"))
    starts.sort()
    colmap = {}  # col idx -> (soort, stat)
    for gi, (start, soort) in enumerate(starts):
        end = starts[gi + 1][0] if gi + 1 < len(starts) else len(head_row)
        for i in range(start, end):
            h = head_row[i]
            if h is None:
                continue
            key = TEXT_KEYS.get(norm(h))
            if key:
                colmap[i] = (soort, key)

    totals = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        name = r[0] if r else None
        if name is None or str(name).strip() == "":
            continue
        if is_total(name):
            break
        roster.add(name, "Selectie")
        t = {"comp": dict.fromkeys(STAT_KEYS, 0), "oefen": dict.fromkeys(STAT_KEYS, 0)}
        for i, (soort, key) in colmap.items():
            v = num(r[i]) if i < len(r) else None
            if v is not None:
                t[soort][key] = int(v)
        totals[norm(name)] = t
    return totals


# ---------------------------------------------------------------- wedstrijden

SCORE_RE = re.compile(r"(\d+)\s*[-–]\s*(\d+)\s*(?:\(\s*([WGVwgv])\s*\))?")


def parse_score(raw, home: bool, where: str):
    if raw is None or str(raw).strip() == "":
        return None
    m = SCORE_RE.search(str(raw))
    if not m:
        warn(f"{where}: uitslag '{raw}' niet herkend (verwacht bv. '3-1 (W)').")
        return None
    a, b, letter = int(m.group(1)), int(m.group(2)), (m.group(3) or "").upper()
    if letter == "W":
        ours, theirs = max(a, b), min(a, b)
        if a == b:
            warn(f"{where}: uitslag '{raw}' is gelijk maar staat als W.")
    elif letter == "V":
        ours, theirs = min(a, b), max(a, b)
        if a == b:
            warn(f"{where}: uitslag '{raw}' is gelijk maar staat als V.")
    elif letter == "G":
        ours, theirs = a, b
        if a != b:
            warn(f"{where}: uitslag '{raw}' staat als G maar is geen gelijkspel.")
    else:
        # Geen letter: thuisploeg eerst.
        ours, theirs = (a, b) if home else (b, a)
        warn(f"{where}: uitslag '{raw}' zonder (W/G/V); aangenomen dat de thuisploeg eerst staat.")
    result = "W" if ours > theirs else "V" if ours < theirs else "G"
    return {"ours": ours, "theirs": theirs, "result": result}


ITEM_SPLIT = re.compile(r"\s*(?:,|;|&|\ben\b)\s*")


def parse_item(item: str):
    item = item.strip().strip(".").strip() if item.strip().lower() not in OWN_GOAL else item.strip()
    for pat in (r"^(.*?)\s*\(?\s*(\d+)\s*x\s*\)?$", r"^(.*?)\s*\(?\s*x\s*(\d+)\s*\)?$"):
        m = re.match(pat, item, re.I)
        if m and m.group(1).strip():
            return m.group(1).strip(), int(m.group(2))
    m = re.match(r"^(\d+)\s*x\s+(.+)$", item, re.I)
    if m:
        return m.group(2).strip(), int(m.group(1))
    return item, 1


def parse_stats_text(text, roster: Roster, where: str):
    events = {k: [] for k in STAT_KEYS}
    if not text:
        return events
    for line in str(text).splitlines():
        line = line.strip()
        if not line:
            continue
        if ":" not in line:
            warn(f"{where}: regel '{line}' in Stats niet herkend (verwacht 'Doelpunten: ...').")
            continue
        k, v = line.split(":", 1)
        key = TEXT_KEYS.get(norm(k))
        if not key:
            warn(f"{where}: onbekend kopje '{k.strip()}' in Stats (gebruik Doelpunten/Assists/Snuiter/Geel/Rood).")
            continue
        for raw in ITEM_SPLIT.split(v.strip()):
            if not raw.strip() or raw.strip() in {"-", "geen", "Geen"}:
                continue
            if raw.strip().lower() in OWN_GOAL:
                events[key].append({"id": None, "name": "Eigen goal", "n": 1, "og": True})
                continue
            name, n = parse_item(raw)
            p, how = roster.resolve(name)
            if p:
                if how == "voornaam":
                    notes.append(f"{where}: '{name}' gekoppeld aan {p['name']}.")
                events[key].append({"id": p["id"], "name": p["name"], "n": n})
            else:
                extra = " (meerdere spelers met die voornaam)" if how == "dubbel" else " (leenspeler? telt niet mee in spelersstats)"
                warn(f"{where}: '{name}' bij {k.strip()} niet gevonden in de selectie{extra}.")
                events[key].append({"id": None, "name": name, "n": n})
    return events


def read_matches(ws, roster: Roster, season: Season):
    header = [norm(c.value) if c.value else "" for c in ws[1]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    c = {
        "date": col("datum"), "soort": col("soort"), "opp": col("tegenstander"),
        "waar": col("waar", "thuis/uit"), "score": col("uitslag"),
        "stats": col("stats"), "notes": col("bijzonderheden", "opmerkingen"),
    }
    if c["date"] is None or c["opp"] is None:
        sys.exit("FOUT: tabblad Wedstrijden mist kolom 'Datum' of 'Tegenstander'.")

    def cell(r, key):
        i = c[key]
        return r[i] if i is not None and i < len(r) else None

    matches = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = as_date(cell(r, "date"))
        if not d:
            continue
        d = season.fix(d, "Wedstrijden")
        opp = str(cell(r, "opp") or "").strip()
        where = f"Wedstrijd {d:%d-%m} {opp}"
        waar = norm(cell(r, "waar") or "")
        home = waar.startswith("thuis")
        if not waar:
            warn(f"{where}: geen Thuis/Uit ingevuld.")
        soort_raw = norm(cell(r, "soort") or "")
        soort = "oefen" if soort_raw.startswith("oefen") else "comp" if soort_raw.startswith("comp") or soort_raw == "" else None
        if soort is None:
            warn(f"{where}: soort '{cell(r, 'soort')}' onbekend, behandeld als competitie.")
            soort = "comp"
        score = parse_score(cell(r, "score"), home, where)
        if score and not soort_raw:
            warn(f"{where}: gespeeld maar geen Soort ingevuld, behandeld als competitie.")
        events = parse_stats_text(cell(r, "stats"), roster, where)
        notes_txt = cell(r, "notes")
        m = {
            "id": f"{d:%Y-%m-%d}-{slug(opp)}",
            "date": d.strftime("%Y-%m-%d"),
            "time": d.strftime("%H:%M") if (d.hour or d.minute) else None,
            "soort": soort,
            "opponent": opp,
            "home": home,
            "played": score is not None,
            **(score or {}),
            **{k: v for k, v in events.items()},
            "notes": str(notes_txt).strip() if notes_txt else None,
            "lineup": [],
        }
        if score:
            scored = sum(e["n"] for e in events["goals"])
            if events["goals"] and scored != score["ours"]:
                warn(f"{where}: {scored} doelpunten in Stats, maar uitslag zegt {score['ours']} voor ons.")
            elif not events["goals"] and score["ours"] > 0:
                warn(f"{where}: geen doelpuntenmakers ingevuld bij Stats.")
        elif any(events[k] for k in STAT_KEYS):
            warn(f"{where}: Stats ingevuld maar geen uitslag.")
        matches.append(m)
    matches.sort(key=lambda m: (m["date"], m["time"] or ""))
    return matches


# ---------------------------------------------------------------- competitiestand

def team_key(name) -> str:
    """'Merino's de 2' en 'De Merino's 2' -> dezelfde sleutel."""
    return " ".join(sorted(re.sub(r"[^a-z0-9']+", " ", norm(name)).split()))


def read_stand(matches, ours):
    """Leest data/stand.json (overgenomen van de clubsite) en controleert hem."""
    if not STAND_JSON.exists():
        notes.append("Geen data/stand.json gevonden: de site toont geen competitiestand.")
        return None
    st = json.loads(STAND_JSON.read_text(encoding="utf-8"))
    rows = st.get("rows") or []
    for r in rows:
        r["key"] = team_key(r["team"])
        r["us"] = r["key"] == team_key(OUR_TEAM)
    keys = {r["key"] for r in rows}

    us = next((r for r in rows if r["us"]), None)
    if not us:
        warn(f"Stand: '{OUR_TEAM}' niet gevonden in data/stand.json.")
    else:
        mine = {"played": ours["played"], "w": ours["w"], "d": ours["g"], "l": ours["v"],
                "gf": ours["gf"], "ga": ours["ga"], "pts": ours["points"]}
        if us["played"] < mine["played"]:
            warn(f"Stand is verouderd: daarin {us['played']} gespeeld, in de Excel {mine['played']}. "
                 "Haal de stand opnieuw op (zie CLAUDE.md).")
        elif us["played"] > mine["played"]:
            warn(f"Stand is verder dan de Excel: daarin {us['played']} gespeeld, in de Excel {mine['played']}. "
                 "Mist er een wedstrijd in de Excel?")
        else:
            diff = [f"{k} {us[k]} vs {mine[k]}" for k in mine if us[k] != mine[k]]
            if diff:
                warn("Stand en Excel verschillen voor " + OUR_TEAM + ": " + ", ".join(diff) + " (stand vs Excel).")

    for m in matches:
        m["oppKey"] = team_key(m["opponent"])
        if m["soort"] == "comp" and rows and m["oppKey"] not in keys:
            warn(f"Tegenstander '{m['opponent']}' niet gevonden in de stand (andere schrijfwijze?).")
    return {"competitie": st.get("competitie"), "bron": st.get("bron"),
            "opgehaald": st.get("opgehaald"), "rows": rows}


# ---------------------------------------------------------------- main

def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        sys.exit(f"FOUT: {xlsx} bestaat niet.")
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    ws_train = sheet(wb, "Training")
    ws_comp = sheet(wb, "Competitie")
    ws_oefen = sheet(wb, "Oefen")
    ws_match = sheet(wb, "Wedstrijden")
    ws_sel = sheet(wb, "Selectie")

    # Seizoen bepalen op basis van alle datums
    all_dates = []
    for ws in (ws_train, ws_comp, ws_oefen):
        all_dates += [as_date(c.value) for c in ws[1] if as_date(c.value)]
    for r in ws_match.iter_rows(min_row=2, max_col=1, values_only=True):
        if as_date(r[0]):
            all_dates.append(as_date(r[0]))
    if not all_dates:
        sys.exit("FOUT: geen datums gevonden.")
    season = Season(all_dates)

    roster = Roster()
    totals = read_selectie(ws_sel, roster)
    train_sessions, train_rows = read_grid(ws_train, roster, season, "Training")
    comp_cols, comp_rows = read_grid(ws_comp, roster, season, "Competitie")
    oefen_cols, oefen_rows = read_grid(ws_oefen, roster, season, "Oefen")
    matches = read_matches(ws_match, roster, season)

    # Ontbrekende spelers per tabblad
    grids = {"Selectie": totals, "Training": train_rows, "Competitie": comp_rows, "Oefen": oefen_rows}
    for key in roster.order:
        p = roster.players[key]
        missing = [s for s, rows in grids.items() if key not in rows]
        if missing:
            warn(f"{p['name']} ontbreekt in tabblad {', '.join(missing)} (telt daar als 0).")

    # Minuten per wedstrijd koppelen via datum
    by_date = {}
    for soort, cols, rows in (("comp", comp_cols, comp_rows), ("oefen", oefen_cols, oefen_rows)):
        for s in cols:
            by_date[(soort, s["date"].strftime("%Y-%m-%d"))] = (s["col"], rows)
    used = set()
    for m in matches:
        k = (m["soort"], m["date"])
        if k in by_date:
            used.add(k)
            idx, rows = by_date[k]
            lineup = []
            for key, vals in rows.items():
                v = vals.get(idx)
                if v and v > 0:
                    lineup.append({"id": roster.players[key]["id"], "min": int(v)})
            lineup.sort(key=lambda x: (-x["min"], x["id"]))
            m["lineup"] = lineup
        elif m["played"]:
            tab = "Competitie" if m["soort"] == "comp" else "Oefen"
            warn(f"Wedstrijd {m['date']} {m['opponent']}: geen minuten gevonden in tabblad {tab} voor deze datum.")
    for (soort, date) in by_date:
        if (soort, date) not in used:
            tab = "Competitie" if soort == "comp" else "Oefen"
            warn(f"Tabblad {tab} heeft minuten op {date}, maar er staat geen {tab.lower()}wedstrijd op die datum in Wedstrijden.")

    # Spelersstats
    played = [m for m in matches if m["played"]]
    id_to_key = {p["id"]: k for k, p in roster.players.items()}
    players_out = []
    for key in sorted(roster.order, key=lambda k: roster.players[k]["name"].lower()):
        p = roster.players[key]
        out = {"id": p["id"], "name": p["name"]}
        for soort in ("comp", "oefen"):
            ms = [m for m in played if m["soort"] == soort]
            mins = [x["min"] for m in ms for x in m["lineup"] if x["id"] == p["id"]]
            t = totals.get(key, {}).get(soort, dict.fromkeys(STAT_KEYS, 0))
            out[soort] = {"apps": len(mins), "min": sum(mins), **t}
        # Training
        rows = train_rows.get(key)
        if rows is not None:
            att = [1 if (rows.get(s["col"]) or 0) > 0 else 0 for s in train_sessions]
            streak = 0
            for a in reversed(att):
                if not a:
                    break
                streak += 1
            out["train"] = {"att": sum(att), "held": len(att), "streak": streak}
        else:
            out["train"] = {"att": 0, "held": len(train_sessions), "streak": 0}
        players_out.append(out)

    # Controle: handmatige totalen (Selectie) vs Stats-tekst (Wedstrijden)
    labels = {"goals": "goals", "assists": "assists", "snuiter": "snuiters", "geel": "gele kaarten", "rood": "rode kaarten"}
    for soort, soort_label in (("comp", "competitie"), ("oefen", "oefen")):
        ms = [m for m in played if m["soort"] == soort]
        for stat in STAT_KEYS:
            text_has_any = any(m[stat] for m in ms)
            if stat in ("geel", "rood") and not text_has_any:
                continue  # kaarten worden (nog) niet per wedstrijd bijgehouden
            from_text = {}
            for m in ms:
                for e in m[stat]:
                    if e["id"]:
                        from_text[e["id"]] = from_text.get(e["id"], 0) + e["n"]
            for po in players_out:
                a = po[soort][stat]
                b = from_text.get(po["id"], 0)
                if a != b:
                    warn(f"{po['name']}: {a} {labels[stat]} ({soort_label}) in Selectie, maar {b} volgens de wedstrijdverslagen.")

    # Teamstats
    team = {}
    for soort in ("comp", "oefen"):
        ms = [m for m in played if m["soort"] == soort]
        w = sum(m["result"] == "W" for m in ms)
        g = sum(m["result"] == "G" for m in ms)
        v = sum(m["result"] == "V" for m in ms)
        team[soort] = {
            "played": len(ms), "w": w, "g": g, "v": v,
            "gf": sum(m["ours"] for m in ms), "ga": sum(m["theirs"] for m in ms),
            "points": 3 * w + g,
        }

    trainings = []
    for s in train_sessions:
        present = [roster.players[k]["id"] for k, vals in train_rows.items() if (vals.get(s["col"]) or 0) > 0]
        trainings.append({"date": s["date"].strftime("%Y-%m-%d"), "present": sorted(present)})

    stand = read_stand(matches, team["comp"])

    now = dt.datetime.now()
    data = {
        "generated": now.strftime("%Y-%m-%dT%H:%M"),
        "season": season.label,
        "team": team,
        "players": players_out,
        "matches": matches,
        "trainings": trainings,
        "stand": stand,
    }

    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    OUT_JS.parent.mkdir(parents=True, exist_ok=True)
    OUT_JS.write_text(f"/* Gegenereerd door build.py op {data['generated']} - niet met de hand aanpassen */\nwindow.WB3_DATA = {payload};\n", encoding="utf-8")

    # Cache-buster in de HTML zodat bezoekers meteen de nieuwe versie zien
    version = hashlib.sha1(payload.encode()).hexdigest()[:8]

    def bust(m):
        path = SITE / m.group(1)
        h = hashlib.sha1(path.read_bytes()).hexdigest()[:8] if path.exists() else version
        return f"{m.group(1)}?v={h}"

    for html in SITE.glob("*.html"):
        src = html.read_text(encoding="utf-8")
        new = re.sub(r"(assets/[\w.-]+\.(?:js|css))(?:\?v=\w+)?", bust, src)
        if new != src:
            html.write_text(new, encoding="utf-8")

    # Rapport
    tc, to = team["comp"], team["oefen"]
    print(f"Seizoen {season.label} - {xlsx.name}")
    print(f"  {len(players_out)} spelers")
    print(f"  Competitie: {tc['played']} gespeeld ({tc['w']}W {tc['g']}G {tc['v']}V, {tc['gf']}-{tc['ga']}, {tc['points']} pnt)")
    print(f"  Oefen:      {to['played']} gespeeld ({to['w']}W {to['g']}G {to['v']}V, {to['gf']}-{to['ga']})")
    print(f"  Programma:  {sum(not m['played'] for m in matches)} wedstrijden nog te spelen")
    print(f"  Training:   {len(trainings)} trainingen")
    if stand and any(r["us"] for r in stand["rows"]):
        u = next(r for r in stand["rows"] if r["us"])
        print(f"  Stand:      {u['pos']}e van {len(stand['rows'])} ({stand['competitie']}, opgehaald {stand['opgehaald']})")
    if notes:
        print(f"\nOpmerkingen ({len(notes)}):")
        for n in notes:
            print(f"  - {n}")
    if warnings:
        print(f"\nCONTROLEREN ({len(warnings)}):")
        for w in warnings:
            print(f"  ! {w}")
    else:
        print("\nAlle controles OK.")
    print(f"\nGeschreven: {OUT_JS.relative_to(ROOT)} (versie {version})")


if __name__ == "__main__":
    main()
